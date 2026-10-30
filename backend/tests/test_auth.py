"""Sign-in, roles and user management."""
from fastapi.testclient import TestClient

from app.main import app
from app.services.auth_service import hash_password, verify_password
from tests.conftest import ADMIN_EMAIL, ADMIN_PASSWORD, sign_in


def test_password_hashing():
    h = hash_password("s3cret-pass", iterations=1000)
    assert h.startswith("pbkdf2_sha256$1000$") and "s3cret-pass" not in h
    assert verify_password("s3cret-pass", h) and not verify_password("wrong", h)
    assert not verify_password("x", "garbage")


def test_existing_hash_format_is_accepted():
    # same format as the users.json already in Dropbox
    h = hash_password("hello-world")
    algo, iterations, salt, digest = h.split("$")
    assert (algo, iterations) == ("pbkdf2_sha256", "390000") and len(salt) == 24 and len(digest) == 44


def test_data_requires_login(anon):
    assert anon.get("/api/health").status_code == 200          # open
    r = anon.get("/api/sales")
    assert r.status_code == 401 and r.json()["error"]["code"] == "UNAUTHORIZED"
    anon.headers["Authorization"] = "Bearer not-a-token"
    assert anon.get("/api/meta/filters").status_code == 401


def test_default_admin_login_and_me(anon):
    r = anon.post("/api/auth/login", json={"email": ADMIN_EMAIL.upper(), "password": ADMIN_PASSWORD})
    body = r.json()["data"]
    assert r.status_code == 200 and body["token"] and body["user"]["role"] == "admin"
    assert "passwordHash" not in body["user"] and "password_hash" not in body["user"]
    anon.headers["Authorization"] = f"Bearer {body['token']}"
    me = anon.get("/api/auth/me").json()["data"]
    assert me["email"] == ADMIN_EMAIL and me["lastLoginAt"]
    assert anon.get("/api/sales").status_code == 200


def test_wrong_password(anon):
    r = anon.post("/api/auth/login", json={"email": ADMIN_EMAIL, "password": "nope-nope"})
    assert r.status_code == 401 and r.json()["error"]["message"] == "Wrong email or password."


def test_admin_manages_users_and_roles_are_enforced(client, settings, repo):
    r = client.post("/api/users", json={"email": "viewer@test.local", "name": "Vee", "role": "viewer",
                                        "password": "viewer-pass-1"})
    assert r.status_code == 201
    vid = r.json()["data"]["id"]
    assert client.post("/api/users", json={"email": "viewer@test.local", "password": "another-pass"}).status_code == 409
    assert client.post("/api/users", json={"email": "x@test.local", "password": "short"}).status_code == 400
    assert client.post("/api/users", json={"email": "y@test.local", "role": "boss", "password": "long-enough"}).status_code == 400
    emails = [u["email"] for u in client.get("/api/users").json()["data"]]
    assert set(emails) == {ADMIN_EMAIL, "viewer@test.local"}

    with TestClient(app) as v:
        sign_in(v, "viewer@test.local", "viewer-pass-1")
        assert v.get("/api/sales").status_code == 200                     # viewers read
        r = v.post("/api/benefits", json={"name": "X", "tier": "Gold"})
        assert r.status_code == 403 and r.json()["error"]["code"] == "FORBIDDEN"   # but do not change data
        assert v.get("/api/users").status_code == 403                      # and do not manage users

        # promote to editor: data changes allowed now
        assert client.put(f"/api/users/{vid}", json={"role": "editor"}).json()["data"]["role"] == "editor"
        assert v.post("/api/benefits", json={"name": "X", "tier": "Gold"}).status_code != 403

        # deactivate: the token stops working at once
        client.put(f"/api/users/{vid}", json={"active": False})
        assert v.get("/api/sales").status_code == 401
    assert client.delete(f"/api/users/{vid}").status_code == 204
    assert client.get(f"/api/users").json()["data"][0]["email"] == ADMIN_EMAIL


def test_admin_cannot_lock_themselves_out(client):
    me = client.get("/api/auth/me").json()["data"]
    for change in ({"role": "viewer"}, {"active": False}):
        r = client.put(f"/api/users/{me['id']}", json=change)
        assert r.status_code == 403 and r.json()["error"]["code"] == "SELF_LOCKOUT"
    assert client.delete(f"/api/users/{me['id']}").status_code == 403


def test_change_own_password(client):
    r = client.post("/api/auth/change-password", json={"currentPassword": "wrong-one", "newPassword": "new-pass-456"})
    assert r.status_code == 403
    r = client.post("/api/auth/change-password", json={"currentPassword": ADMIN_PASSWORD, "newPassword": "new-pass-456"})
    assert r.status_code == 200
    new_token = r.json()["data"]["token"]
    assert client.get("/api/auth/me").status_code == 401                   # the old token no longer works
    client.headers["Authorization"] = f"Bearer {new_token}"
    assert client.get("/api/auth/me").status_code == 200
    assert client.post("/api/auth/login", json={"email": ADMIN_EMAIL, "password": "new-pass-456"}).status_code == 200


def test_user_register_excel_has_roles_and_history_but_no_passwords(client, repo):
    import io

    import openpyxl

    from app.services import user_register
    r = client.post("/api/users", json={"email": "ed@test.local", "name": "Eddie", "role": "editor",
                                        "password": "editor-secret-1"})
    uid = r.json()["data"]["id"]
    assert client.put(f"/api/users/{uid}", json={"role": "viewer", "password": "new-secret-22"}).status_code == 200
    assert client.delete(f"/api/users/{uid}").status_code == 204

    saved = repo.store.read(user_register.REGISTER_PATH)                        # kept current in the store
    res = client.get("/api/users/register")
    assert res.status_code == 200 and saved
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Users", "History", "Roles"]
    users = [[c.value for c in row] for row in wb["Users"].iter_rows(min_row=5)]
    assert all(u[1] != "ed@test.local" for u in users) and any(u[2] == "admin" for u in users)
    history = [[c.value for c in row] for row in wb["History"].iter_rows(min_row=5)]
    acts = [(h[2], h[3]) for h in history if h[3] == "ed@test.local"]
    assert acts == [("Deleted", "ed@test.local"), ("Changed", "ed@test.local"), ("Created", "ed@test.local")]
    changed = next(h for h in history if h[2] == "Changed")
    assert "role editor → viewer" in changed[6] and "password reset" in changed[6]
    cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "editor-secret-1" not in cells and "new-secret-22" not in cells and "pbkdf2" not in cells
