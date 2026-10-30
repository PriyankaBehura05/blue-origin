"""HTTP-level tests: envelopes, camelCase contract, CRUD round-trips through the store, uploads, reports."""
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import RETAILERS_CSV


def test_health_reports_store(client):
    body = client.get("/api/health").json()
    assert body["status"] == "ok" and body["store_ok"] is True and body["store"].startswith("local:")


def test_cors_allows_vite_dev_server(client):
    r = client.options("/api/health", headers={"Origin": "http://localhost:5173",
                                               "Access-Control-Request-Method": "GET"})
    assert r.headers.get("access-control-allow-origin") == "http://localhost:5173"


def test_filters_endpoint_is_camel_case(client):
    data = client.get("/api/meta/filters").json()["data"]
    assert data["currentQuarter"] == "Q4 FY 26" and data["previousQuarter"] == "Q3 FY 26"
    assert {"id", "name"} == set(data["distributors"][0])
    assert "fulfillmentStatuses" in data and "campaignRegions" not in data


def test_sales_paging_and_aliases(client):
    data = client.get("/api/sales?quarter=Q4 FY 26&page_size=2&sort_by=netAmount&sort_dir=desc").json()["data"]
    assert data["total"] == 3 and data["pageSize"] == 2 and data["page"] == 1
    assert data["items"][0]["netAmount"] == 160000 and data["items"][0]["retailerName"] == "Test Store One"
    assert data["summary"]["net"] == 191000
    did = data["items"][0]["distributorId"]
    assert client.get(f"/api/sales?distributorId={did}").json()["data"]["total"] == 2


def test_bad_params(client):
    r = client.get("/api/sales?page=zero")
    assert r.status_code == 422 and r.json()["error"]["code"] == "REQUEST_VALIDATION_ERROR"
    r = client.get("/api/sales?sort_by=bogus")
    assert r.status_code == 400 and r.json()["error"]["code"] == "INVALID_FILTER"


def test_retailers_and_detail(client):
    page = client.get("/api/retailers?tier=Green").json()["data"]
    assert page["total"] == 1
    r = page["items"][0]
    assert r["currentTier"] == "Green" and r["previousTier"] == "Yellow" and r["inMaster"] is True
    detail = client.get(f"/api/retailers/{r['id']}").json()["data"]
    assert detail["retailer"]["name"] == "Test Store One" and len(detail["sales"]) == 2
    assert client.get("/api/retailers/unknown-id").json()["error"]["code"] == "NOT_FOUND"


def test_distributors_tiers(client):
    assert client.get("/api/distributors?region=South").json()["data"]["total"] == 2
    defs = client.get("/api/tiers/definitions").json()["data"]
    assert defs[0]["tier"] == "Unranked" and "minQuarterlySales" in defs[0]
    moves = client.get("/api/tiers/movements?direction=Upgraded").json()["data"]["items"]
    assert moves[0]["from"] == "Yellow" and moves[0]["to"] == "Green"
    assert client.get("/api/tiers/classify?amount=200000").json()["data"]["tier_code"] == "GREEN"


def test_campaigns_are_gone(client):
    assert client.get("/api/campaigns").status_code == 404


def test_benefit_and_fulfilment_crud(client):
    r = client.post("/api/benefits", json={"name": "Gift Hamper", "tier": "gold", "frequency": "Annual",
                                           "quantity": 2, "unitCost": 4000})
    assert r.status_code == 201
    ben = r.json()["data"]
    assert ben["tier"] == "Gold" and ben["id"] == "BEN-0003"
    assert client.post("/api/benefits", json={"name": "Bad", "tier": "Diamond"}).status_code == 400

    rid = client.get("/api/retailers?search=Test Store Two").json()["data"]["items"][0]["id"]
    r = client.post("/api/benefits/fulfillments", json={"benefitId": "BEN-0003", "retailerId": rid,
                                                        "plannedDate": "2099-01-15"})
    assert r.status_code == 201
    ful = r.json()["data"]
    assert ful["status"] == "Planned" and ful["cost"] == 8000 and ful["retailerName"] == "Test Store Two"

    r = client.put(f"/api/benefits/fulfillments/{ful['id']}",
                   json={"benefitId": "BEN-0003", "retailerId": rid, "plannedDate": "2099-01-15",
                         "actualDate": "2026-10-01", "status": "Fulfilled"})
    assert r.json()["data"]["status"] == "Fulfilled"

    in_use = client.delete("/api/benefits/BEN-0003")
    assert in_use.status_code == 400 and in_use.json()["error"]["code"] == "IN_USE"
    assert client.delete(f"/api/benefits/fulfillments/{ful['id']}").status_code == 204
    assert client.delete("/api/benefits/BEN-0003").status_code == 204


def test_fulfilment_list_summary(client):
    data = client.get("/api/benefits/fulfillments?status=Overdue").json()["data"]
    assert data["total"] == 1 and data["summary"]["overdue"] == 1


def test_reports_preview_and_download_logs_generation(client):
    defs = client.get("/api/reports").json()["data"]
    assert len(defs) == 8 and defs[0]["lastGenerated"] is None
    preview = client.get("/api/reports/RPT-03/preview").json()["data"]
    assert preview["reportId"] == "RPT-03" and preview["totalRows"] == 2

    r = client.get("/api/reports/RPT-03/download?format=xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    ws = load_workbook(BytesIO(r.content)).active
    # TechnoSport report format: title, notes, blank, header (row 4), 2 distributors, TOTAL row
    assert ws["A1"].value == "TECHNOSPORT  |  Distributor Performance" and ws["A1"].font.name == "Arial"
    assert ws.cell(4, 1).value == "Distributor" and ws.cell(4, 1).fill.fgColor.rgb == "FFF2F2F2"
    assert ws.max_row == 7 and ws.cell(7, 1).value == "TOTAL" and ws.cell(7, 5).value == "=SUM(E5:E6)"
    assert ws.cell(5, 5).number_format == '###,##0;\(###,##0\);"–"' and ws.freeze_panes == "B5"
    assert client.get("/api/reports").json()["data"][2]["lastGenerated"] is not None

    csv = client.get("/api/reports/RPT-04/download?format=csv")
    assert csv.headers["content-type"].startswith("text/csv") and "Test Store One" in csv.content.decode("utf-8-sig")


def test_datasets_status_template_and_download(client):
    data = client.get("/api/datasets").json()["data"]
    sales = next(d for d in data if d["dataset"] == "sales")
    assert sales["stored"] and sales["rows"] == 6 and sales["validRows"] == 4
    assert "net_sales_amount" in sales["requiredColumns"]
    assert client.get("/api/datasets/targets/template").text.startswith("month,target_amount")
    assert "Test Dist A" in client.get("/api/datasets/distributors/download").text


def test_upload_validate_csv(client):
    files = {"file": ("retailers.csv", RETAILERS_CSV.encode(), "text/csv")}
    report = client.post("/api/uploads/validate", files=files, data={"dataset": "retailers"}).json()["data"]
    assert report["can_process"] and report["total_rows"] == 3
    assert report["status_counts"]["valid"] == 3 and len(report["preview"]) == 3


def test_upload_validate_missing_columns_reported(client):
    files = {"file": ("x.csv", b"foo,bar\n1,2\n", "text/csv")}
    report = client.post("/api/uploads/validate", files=files, data={"dataset": "sales"}).json()["data"]
    assert not report["can_process"]
    assert "net_sales_amount" in report["missing_required_columns"]


def test_upload_rejects_empty_bad_type_and_oversize(client):
    assert client.post("/api/uploads/validate", files={"file": ("e.csv", b"", "text/csv")},
                       data={"dataset": "sales"}).json()["error"]["code"] == "EMPTY_FILE"
    assert client.post("/api/uploads/validate", files={"file": ("e.pdf", b"%PDF", "application/pdf")},
                       data={"dataset": "sales"}).json()["error"]["code"] == "INVALID_FILE"
    big = b"a\n" + b"x" * (1024 * 1024 + 10)
    r = client.post("/api/uploads/validate", files={"file": ("big.csv", big, "text/csv")}, data={"dataset": "sales"})
    assert r.status_code == 413


def test_upload_store_history_and_duplicate(client, store):
    files = {"file": ("retailers.csv", RETAILERS_CSV.encode(), "text/csv")}
    r1 = client.post("/api/uploads", files=files, data={"dataset": "retailers"})
    assert r1.status_code == 201
    stored = r1.json()["data"]["upload"]["stored_filename"]
    assert store.read(f"uploads/{stored}") == RETAILERS_CSV.encode()

    r2 = client.post("/api/uploads", files=files, data={"dataset": "retailers"})
    assert r2.status_code == 409 and r2.json()["error"]["code"] == "DUPLICATE_UPLOAD"

    r3 = client.post("/api/uploads?allow_duplicate=true", files=files, data={"dataset": "retailers"})
    assert r3.status_code == 201
    assert len(client.get("/api/uploads").json()["data"]) == 2


def test_upload_activate_replace_and_append(client):
    new_rows = "Retailer Name,Distributor,Retailer Code,Status\nBrand New Store,Test Dist A,R9,Active\n"
    files = {"file": ("more.csv", new_rows.encode(), "text/csv")}
    up = client.post("/api/uploads", files=files, data={"dataset": "retailers"}).json()["data"]["upload"]

    r = client.post(f"/api/uploads/{up['upload_id']}/activate?mode=append")
    assert r.status_code == 200 and r.json()["data"]["rows_stored"] == 4
    names = {x["name"] for x in client.get("/api/retailers?page_size=50").json()["data"]["items"]}
    assert "Brand New Store" in names and "Test Store One" in names

    r = client.post(f"/api/uploads/{up['upload_id']}/activate?mode=replace")
    assert r.json()["data"]["rows_stored"] == 1
    hist = client.get("/api/uploads").json()["data"][0]
    assert hist["processed"] is True and hist["activation_mode"] == "replace"
    master = [x for x in client.get("/api/retailers?page_size=50").json()["data"]["items"] if x["inMaster"]]
    assert [x["name"] for x in master] == ["Brand New Store"]


def test_activate_rejects_unprocessable_upload(client):
    files = {"file": ("x.csv", b"foo,bar\n1,2\n", "text/csv")}
    up = client.post("/api/uploads", files=files, data={"dataset": "sales"}).json()["data"]["upload"]
    r = client.post(f"/api/uploads/{up['upload_id']}/activate")
    assert r.status_code == 400 and r.json()["error"]["code"] == "UPLOAD_NOT_PROCESSABLE"


def test_unknown_dataset_type_rejected(client):
    files = {"file": ("r.csv", RETAILERS_CSV.encode(), "text/csv")}
    r = client.post("/api/uploads/validate", files=files, data={"dataset": "invoices"})
    assert r.status_code == 422


def test_empty_store_serves_empty_pages(tmp_path, settings):
    """A brand-new Dropbox folder: every page loads, with no data (the first admin is created on sign-in)."""
    from fastapi.testclient import TestClient

    from app.api.deps import get_repository, get_settings
    from app.main import app
    from app.repositories.store_repository import StoreRepository
    from app.storage.local_store import LocalStore
    from tests.conftest import ADMIN_EMAIL, ADMIN_PASSWORD, sign_in

    repo = StoreRepository(LocalStore(tmp_path / "empty"))
    app.dependency_overrides[get_repository] = lambda: repo
    app.dependency_overrides[get_settings] = lambda: settings
    try:
        with TestClient(app) as c:
            sign_in(c, ADMIN_EMAIL, ADMIN_PASSWORD)
            assert c.get("/api/network/data").status_code == 503
            assert c.get("/api/meta/filters").status_code == 200
            for path in ("/api/sales", "/api/retailers", "/api/distributors", "/api/tiers/movements", "/api/benefits",
                         "/api/benefits/fulfillments"):
                assert c.get(path).json()["data"]["total"] == 0, path
            assert c.get("/api/reports/RPT-01/preview").json()["data"]["rows"] == []
    finally:
        app.dependency_overrides.clear()


def test_tiers_for_any_quarter(client):
    filters = client.get("/api/meta/filters").json()["data"]
    assert filters["salesQuarters"] == ["Q4 FY 26", "Q3 FY 26"]
    assert client.get("/api/tiers/definitions?quarter=bogus").status_code == 400

    defs = {d["tier"]: d["retailerCount"] for d in client.get("/api/tiers/definitions?quarter=Q3 FY 26").json()["data"]}
    assert defs["Yellow"] == 1 and defs["Green"] == 0
    rows = client.get("/api/tiers/movements?quarter=Q4 FY 26&search=Test Dist A").json()["data"]["items"]
    assert rows[0]["to"] == "Green" and rows[0]["previousQuarterSales"] == 60000 and rows[0]["distributorName"] == "Test Dist A"
    mix = client.get("/api/tiers/by-quarter").json()["data"]
    assert [r["quarter"] for r in mix] == ["Q4 FY 26", "Q3 FY 26"]
    assert mix[0]["counts"]["Green"] == 1 and mix[0]["buyingRetailers"] == 3


def test_retailer_and_distributor_exports(client):
    r = client.get("/api/retailers/export?format=xlsx&tier=Green")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    ws = load_workbook(BytesIO(r.content)).active
    assert ws.cell(4, 2).value == "Retailer" and ws.max_row == 5 and ws.cell(5, 2).value == "Test Store One"
    assert ws.auto_filter.ref == "A4:P5" and ws.cell(5, 7).value == "560001"     # pincode stays text
    csv = client.get("/api/retailers/export?format=csv&sort_by=name").content.decode("utf-8-sig").splitlines()
    assert len(csv) == 5 and csv[0].startswith("Retailer Code,Retailer")
    d = client.get("/api/distributors/export?format=csv").content.decode("utf-8-sig").splitlines()
    assert len(d) == 3 and "Test Dist A" in "".join(d)
    assert client.get("/api/retailers/export?format=pdf").status_code == 422


def test_report_filters_apply_to_preview_and_download(client):
    opts = client.get("/api/meta/filters").json()["data"]
    assert opts["months"][:3] == ["2026-03", "2026-02", "2026-01"]
    dist_a = next(d["id"] for d in opts["distributors"] if d["name"] == "Test Dist A")

    def rows(rid, **f):
        q = "&".join(f"{k}={v}" for k, v in f.items())
        return client.get(f"/api/reports/{rid}/preview?{q}").json()["data"]

    by_state = rows("RPT-03", state="Karnataka")
    assert by_state["rows"] == [["Test Dist A", "D1", "South", 1, 160000]] and by_state["filters"] == "Karnataka"
    earlier = rows("RPT-03", quarter="Q3 FY 26")                     # Oct-Dec 2025: only Test Dist A bought
    assert earlier["columns"][-1] == "Q3 FY 26 Net Sales (₹)" and earlier["rows"][0][-1] == 60000
    assert rows("RPT-02", month="2026-02")["rows"] == [["Feb 26", "South", 30000]]
    jan = rows("RPT-01", month="2026-01")
    assert jan["columns"][-1] == "Jan 26 Net Sales (₹)" and [r[0] for r in jan["rows"]] == ["Test Store One"]
    assert all(r[2] == "Test Dist A" for r in rows("RPT-04", distributor=dist_a)["rows"])

    r = client.get("/api/reports/RPT-03/download?format=csv&state=Karnataka&quarter=Q4%20FY%2026")
    assert "Q4_FY_26_Karnataka" in r.headers["content-disposition"]
    body = r.content.decode("utf-8-sig")
    assert "Test Dist A" in body and "Test Dist B" not in body
    assert client.get("/api/reports/RPT-03/preview?month=2026-13").status_code != 200
