"""
Excel downloads in the TechnoSport report format (the layout of Distributor_Counter_Report.xlsx):

    row 1   TECHNOSPORT  |  <title>                  Arial 12 bold, merged, grey rule underneath
    row 2   notes · filters · generated date         Arial 9 italic grey, merged
    row 3   (blank)
    row 4   header                                   Arial 10 bold on F2F2F2, thin BFBFBF borders
    row 5+  data                                     Arial 10, thin borders; text left, counts centred,
                                                     ₹ right as ###,##0 (negatives in brackets, zero as –),
                                                     % as green / red 0.0%
    last    TOTAL (summary tables)                   bold on E6E6E6, =SUM() formulas
Total / Peak columns are shaded E6E6E6. Panes are frozen under the header and after the first column;
detail tables (no TOTAL row) get header filters.
"""
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, Optional, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
INR_FMT = '###,##0;\\(###,##0\\);"–"'
NUM_FMT = '#,##0;\\-#,##0;"–"'
PCT_FMT = '[Green]0.0%;[Red]\\-0.0%;"–"'
HEADER_FILL = PatternFill("solid", fgColor="FFF2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="FFE6E6E6")
_thin = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
TITLE_RULE = Border(bottom=Side(style="medium", color="FF808080"))
HEADER_ROW = 4

_INR_WORDS = re.compile(r"₹|\b(sales|revenue|amount|cost|value|purchase|spend)\b", re.I)
_TOTAL_COLS = re.compile(r"^(total|fy total|peak|lifetime)\b", re.I)
_ID_COLS = re.compile(r"pin ?code|phone|mobile|\bcode\b|\bid\b|year|month", re.I)  # numbers that are labels


@dataclass
class Sheet:
    title: str
    df: pd.DataFrame
    notes: str = ""
    total: bool = False                      # add a TOTAL row (summary tables)
    kinds: Dict[str, str] = field(default_factory=dict)  # column -> text | num | inr | pct (pct = percent units)


def column_kind(name: str, s: pd.Series) -> str:
    if "%" in str(name):
        return "pct"
    if _ID_COLS.search(str(name)):
        return "text"
    numeric = pd.to_numeric(s, errors="coerce")
    if len(s) and numeric.notna().sum() < s.notna().sum():   # some values are not numbers -> text
        return "text"
    if s.notna().sum() == 0:
        return "text"
    return "inr" if _INR_WORDS.search(str(name)) else "num"


def _value(v, kind: str):
    if v is None or (isinstance(v, float) and v != v) or v is pd.NA:
        return None
    if kind in ("num", "inr", "pct"):
        try:
            n = float(v)
        except (TypeError, ValueError):
            return v
        if kind == "pct":
            return n / 100
        return int(n) if n.is_integer() else round(n, 2)
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def _write_sheet(wb: Workbook, sh: Sheet, used: set) -> None:
    name = re.sub(r"[\\/*?:\[\]]", "-", sh.title)[:31]
    base, k = name, 2
    while name in used:
        name = f"{base[:27]} ({k})"
        k += 1
    used.add(name)
    ws = wb.create_sheet(name)
    df = sh.df
    cols = [str(c) for c in df.columns]
    ncol = max(len(cols), 1)
    kinds = {c: sh.kinds.get(c) or column_kind(c, df[c]) for c in cols}
    last = get_column_letter(ncol)

    ws["A1"] = f"TECHNOSPORT  |  {sh.title}"
    ws["A1"].font = Font(name=FONT, size=12, bold=True, color="FF1A1A1A")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    for c in range(1, ncol + 1):
        ws.cell(1, c).border = TITLE_RULE
    notes = "  •  ".join(x for x in (sh.notes, f"Generated {datetime.now().strftime('%d %b %Y, %H:%M')}") if x)
    ws["A2"] = notes
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="FF777777")
    if ncol > 1:
        ws.merge_cells(f"A1:{last}1")
        ws.merge_cells(f"A2:{last}2")

    head_font = Font(name=FONT, size=10, bold=True, color="FF1A1A1A")
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font = head_font
        cell.fill = TOTAL_FILL if _TOTAL_COLS.match(c) else HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center", wrap_text=True)
    ws.row_dimensions[HEADER_ROW].height = 26.4

    plain, bold = Font(name=FONT, size=10, color="FF000000"), Font(name=FONT, size=10, bold=True, color="FF000000")
    align = {"text": Alignment(horizontal="left", vertical="center"), "num": Alignment(horizontal="center", vertical="center"),
             "inr": Alignment(horizontal="right", vertical="center"), "pct": Alignment(horizontal="center", vertical="center")}
    fmts = {"num": NUM_FMT, "inr": INR_FMT, "pct": PCT_FMT}
    shaded = [_TOTAL_COLS.match(c) is not None for c in cols]
    first = HEADER_ROW + 1
    for i, row in enumerate(df.itertuples(index=False, name=None)):
        r = first + i
        for j, (c, v) in enumerate(zip(cols, row), start=1):
            kind = kinds[c]
            cell = ws.cell(r, j, _value(v, kind))
            cell.border = BORDER
            cell.alignment = align[kind]
            if kind in fmts:
                cell.number_format = fmts[kind]
            if shaded[j - 1]:
                cell.font, cell.fill = bold, TOTAL_FILL
            else:
                cell.font = plain
    end = first + len(df) - 1

    if sh.total and len(df):
        t = end + 1
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(t, j)
            if j == 1:
                cell.value = "TOTAL"
            elif kinds[c] in ("num", "inr") and not re.search(r"\b(rate|code|id|pincode|year)\b", c, re.I):
                col = get_column_letter(j)
                cell.value = f"=SUM({col}{first}:{col}{end})"
                cell.number_format = fmts[kinds[c]]
            cell.font, cell.fill, cell.border = bold, TOTAL_FILL, BORDER
            cell.alignment = align[kinds[c]] if j > 1 else align["text"]
    elif len(df):
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last}{end}"

    ws.freeze_panes = f"B{first}"
    for j, c in enumerate(cols, start=1):
        sample = [len(str(v)) for v in df[c].head(300).tolist() if v is not None]
        width = max([len(c) * 0.9] + sample) + 2
        ws.column_dimensions[get_column_letter(j)].width = max(9, min(width if j > 1 else max(width, 24), 48))


def to_xlsx(sheets: Sequence[Sheet]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for sh in sheets:
        _write_sheet(wb, sh, used)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def single(df: pd.DataFrame, title: str, notes: str = "", total: bool = False,
           kinds: Optional[Dict[str, str]] = None) -> bytes:
    return to_xlsx([Sheet(title, df, notes, total, kinds or {})])
